"""Tests pinning the fixes from the final whole-branch review of the deck."""
import pytest
from pptx import Presentation

from deck import build, charts, copy, measure
from deck.facts import load, pillar_counts
from scripts.io import read_csv

PX = 6350


@pytest.fixture(scope="module")
def f():
    return load()


@pytest.fixture(scope="module")
def slides(f):
    return copy.slides(f)


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    return Presentation(str(build.build(tmp_path_factory.mktemp("deck") / "d.pptx")))


@pytest.fixture(scope="module")
def drawn(f, tmp_path_factory):
    charts.render_all(f, tmp_path_factory.mktemp("charts"))
    return charts.LAST_DATA


def _slide(slides, chart):
    return next(s for s in slides if s.get("chart") == chart)


# Critical 1: slide 14 pillar leaders
def test_pillar_counts_reproduce_brand_pillar_when_nothing_is_excluded():
    answers, mentions = read_csv("data/interim/answers.csv"), read_csv("data/interim/mentions.csv")
    mine = pillar_counts(answers, mentions, exclude=())
    bp = read_csv("data/processed/brand_pillar.csv")
    bp = bp[(bp.engine == "All") & (bp.n_answers_mentioning > 0)]
    ref = {(r.pillar, r.brand): (int(r.n), int(r.n_answers_mentioning)) for r in bp.itertuples()}
    got = {(r.pillar, r.brand): (int(r.n), int(r.named)) for r in mine.itertuples()}
    assert got == ref


def test_pillar_leaders_use_only_unbranded_prompts(f):
    pc = f["pillar_counts"]
    assert pc.groupby("pillar").n.first().to_dict() == {
        "Digital experience": 6, "Fees & rates": 8, "Safety & security": 8, "Service": 6, "Trust": 6}
    lead = f["pillar_leaders"]
    assert set(lead["Digital experience"]) == {"Bank Jago", "SeaBank"}
    assert set(lead["Fees & rates"]) == {"Bank Jago", "Krom Bank", "SeaBank"}
    assert lead["Service"] == ["BCA"]


def test_no_pillar_lead_exceeds_one_answer(f):
    # The slide 14 title says no lead is more than one answer; this guards it against new data.
    for pillar, g in f["pillar_counts"].groupby("pillar"):
        top = sorted(g.named, reverse=True)
        assert top[0] - top[1] <= 1, pillar


def test_pillar_slide_names_every_leader_and_drops_the_apps_note(f, slides):
    s = _slide(slides, "pillar_heatmap")
    body = " ".join(s["body"])
    for leaders in f["pillar_leaders"].values():
        for b in leaders:
            assert b in body
    assert "apps" not in s["notes"]
    assert "lead" in s["title"].lower() and "one answer" in s["title"]


def test_heatmap_shows_row_n_and_every_leader_column(f, drawn):
    h = drawn["pillar_heatmap"]
    assert h["row_n"] == f["pillar_counts"].groupby("pillar").n.first().to_dict()
    assert {b for v in f["pillar_leaders"].values() for b in v} <= set(h["banks"])


# Important 2: slide 20 overlap title
def test_overlap_title_states_both_levels(f, slides):
    o = f["overlap"]
    t = _slide(slides, "overlap")["title"]
    assert f"{o['sites_both']} of {o['sites_total']} sites" in t
    assert f"{o['pages_both']} of {o['pages_total']} pages" in t
    assert "barely" not in t


# Important 3: slide 19 lights each engine's claim
def test_source_mix_lights_the_types_each_claim_is_about(drawn):
    assert drawn["source_mix_lit"] == {"ChatGPT": {"bank_official", "regulator"},
                                       "Gemini": {"fintech_platform", "blog_aggregator", "news_media"}}


# Important 4: slide 24 alias caveat
def test_sov_slide_discloses_the_alias_counting_gap(slides):
    s = _slide(slides, "sov_gap")
    text = " ".join(s["body"])
    assert "blu" in text and "“Jago”" in text


# Important 5: slide 12 reads small gaps as close
def test_vis_title_names_banks_close_to_the_top(f, slides):
    v = f["vis"][~f["vis"].low_n].sort_values("vis", ascending=False)
    top = v.vis.iloc[0]
    t = _slide(slides, "vis_rank")["title"]
    for b in v[(top - v.vis) < 5].brand:
        assert b in t, b


# Important 6: line pitch and slide 2 spacing
def test_line_factor_covers_the_rendered_pitch():
    # LibreOffice draws Work Sans at line_spacing 1.15 with a pitch of 1.379 em (slide 2, 22 pt and 14 pt lines).
    assert measure.LINE >= 1.379


def test_cold_open_prompt_and_translation_do_not_touch(deck):
    boxes = sorted((sh for sh in deck.slides[1].shapes if sh.has_text_frame and sh.text_frame.text.strip()),
                   key=lambda sh: sh.top)
    prompt = next(sh for sh in boxes if sh.text_frame.text.startswith("“"))
    english = boxes[boxes.index(prompt) + 1]
    assert english.top / PX - (prompt.top + prompt.height) / PX >= 20


# Important 7: Zaipad at one level, looked up by domain
def test_page_level_counts_do_not_depend_on_row_order(f):
    shuffled = dict(f, top_pages=f["top_pages"].iloc[::-1].reset_index(drop=True))
    assert copy.numbers(shuffled)["zai_page"] == copy.numbers(f)["zai_page"]


def test_slide_28_counts_sites_not_articles(slides):
    s = next(x for x in slides if x["title"] == "For banks")
    item = next(i for i in s["body"] if "Zaipad" in i)
    assert "sites" in item


# Upgraded minors
def test_how_to_gap_is_hedged(slides):
    body = " ".join(_slide(slides, "intent_named")["body"])
    assert "look like a gap" in body


def test_mirror_sites_lights_the_gemini_only_sites(f, drawn):
    t = f["top_sites"]
    assert drawn["mirror_sites_lit"] == set(t[t.gpt == 0].domain)
    assert drawn["mirror_sites_lit"]


def test_every_data_chart_states_its_n(slides):
    for n, s in enumerate(slides, start=1):
        if s.get("chart") and s["chart"] != "pawc_curve":
            assert "n = " in s["foot"], (n, s["chart"])


def test_sov_chart_lights_the_bank_the_title_is_about(drawn):
    # The slide 24 title leads with Bank Mandiri, the gap that does not depend on alias counting.
    assert drawn["sov_gap_lit"] == {"Bank Mandiri"}
