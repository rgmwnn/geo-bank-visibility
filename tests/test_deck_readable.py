"""Chart text must be readable where it lands: on the slide, at its placed size."""
import pytest
from pptx import Presentation

from deck import build, charts

PX = 6350
MIN_PX = 24  # 12 pt on the 1920 px slide; body text is 28 px, footnotes 18 px


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    return Presentation(str(build.build(tmp_path_factory.mktemp("deck") / "d.pptx")))


def test_every_chart_is_placed_and_named(deck):
    names = {sh.name for s in deck.slides for sh in s.shapes if sh.name.startswith("chart-")}
    assert names == {f"chart-{k}" for k in charts.CHARTS}


def test_smallest_chart_text_is_readable_on_the_slide(deck):
    dpi = 100 * charts.DPI_SCALE * charts.SHRINK
    small = []
    for n, s in enumerate(deck.slides, start=1):
        for sh in s.shapes:
            if not sh.name.startswith("chart-"):
                continue
            key = sh.name[len("chart-"):]
            img_w = sh.image.size[0]
            scale = (sh.width / PX) / img_w
            px = charts.LAST_DATA["min_pt"][key] * dpi / 72 * scale
            if px < MIN_PX:
                small.append((n, key, round(px, 1)))
    assert not small, small


def test_no_chart_label_is_clipped_or_overlaps_another(tmp_path):
    from deck.facts import load
    charts.render_all(load(), tmp_path)
    clipped = {k: v for k, v in charts.LAST_DATA["clipped"].items() if v}
    overlaps = {k: v for k, v in charts.LAST_DATA["overlaps"].items() if v}
    assert not clipped, clipped
    assert not overlaps, overlaps
