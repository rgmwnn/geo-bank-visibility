"""Line-language motifs (deck/design-philosophy.md, Event Horizon): present where the spec puts them, never over text."""
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deck import build, charts, copy
from deck.facts import load

PX = 6350


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    return Presentation(str(build.build(tmp_path_factory.mktemp("deck") / "d.pptx")))


@pytest.fixture(scope="module")
def specs():
    return copy.slides(load())


def _box(sh):
    return sh.left / PX, sh.top / PX, (sh.left + sh.width) / PX, (sh.top + sh.height) / PX


def _texts(slide):
    return [_box(sh) for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def _motifs(slide, prefix="motif"):
    return [sh for sh in slide.shapes if sh.name.startswith(prefix)]


def _crosses(sh, rect, pad=4):
    x0, y0, x1, y1 = rect[0] - pad, rect[1] - pad, rect[2] + pad, rect[3] + pad
    if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.name.startswith("motif-line"):
        bx, by, ex, ey = sh.begin_x / PX, sh.begin_y / PX, sh.end_x / PX, sh.end_y / PX
        for i in range(201):
            t = i / 200
            x, y = bx + (ex - bx) * t, by + (ey - by) * t
            if x0 < x < x1 and y0 < y < y1:
                return True
        return False
    a = _box(sh)
    return a[0] < x1 and a[2] > x0 and a[1] < y1 and a[3] > y0


def test_no_motif_touches_text(deck):
    for n, s in enumerate(deck.slides, start=1):
        texts = _texts(s)
        for m in _motifs(s):
            for t in texts:
                assert not _crosses(m, t), (n, m.name, [round(v) for v in t])


PHOTO_SLIDES = [1, 5, 10, 18, 27]


def test_cover_carries_photo_fade_and_one_amber_rule(deck):
    names = [sh.name for sh in deck.slides[0].shapes]
    assert names.count("photo-slot") == 1 and names.count("ground-fade") == 1
    assert [x for x in names if x.startswith("motif-rule")] == ["motif-rule-cover"]


def test_every_slide_is_navy(deck):
    for n, sl in enumerate(deck.slides, start=1):
        assert str(sl.background.fill.fore_color.rgb) == "0F2747", n
        for sh in sl.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.fill.type == 1:  # solid fill
                assert str(sh.fill.fore_color.rgb) != "000000", (n, sh.name)


def test_photo_slots_sit_behind_fade_and_motifs(deck):
    for n, sl in enumerate(deck.slides, start=1):
        shapes = list(sl.shapes)
        slots = [i for i, sh in enumerate(shapes) if sh.name == "photo-slot"]
        if n not in PHOTO_SLIDES:
            assert not slots, n
            continue
        assert len(slots) == 1, n
        later = [i for i, sh in enumerate(shapes) if sh.name.startswith(("motif", "ground"))
                 or (sh.has_text_frame and sh.text_frame.text.strip())]
        assert slots[0] < min(later), n


def test_photo_page_text_sits_on_the_opaque_fade(deck):
    for n in PHOTO_SLIDES:
        sl = deck.slides[n - 1]
        fade = next(sh for sh in sl.shapes if sh.name == "ground-fade")
        opaque = (fade.left + fade.width) / PX - build.FADE_RAMP
        for t in _texts(sl):
            assert t[2] <= opaque, (n, round(t[2]), round(opaque))


def test_chart_panels_have_four_corner_brackets(deck):
    for n, sl in enumerate(deck.slides, start=1):
        pics = [sh for sh in sl.shapes if sh.name.startswith("chart-")]
        if not pics:
            continue
        brackets = _motifs(sl, "motif-bracket")
        assert {m.name.split("-")[2] for m in brackets} == {"tl", "tr", "bl", "br"}, n
        bx0 = min(min(m.begin_x, m.end_x) for m in brackets) / PX
        by0 = min(min(m.begin_y, m.end_y) for m in brackets) / PX
        bx1 = max(max(m.begin_x, m.end_x) for m in brackets) / PX
        by1 = max(max(m.begin_y, m.end_y) for m in brackets) / PX
        for p in pics:
            x0, y0, x1, y1 = _box(p)
            assert bx0 <= x0 and by0 <= y0 and x1 <= bx1 and y1 <= by1, (n, p.name)


def test_gauges_encode_their_values(deck):
    names = [sh.name for sh in deck.slides[2].shapes]
    assert "motif-gauge-1" in names and "motif-gauge-2" in names
    assert abs(build.GAUGES[1] - 0.24 * 360) <= 1 and abs(build.GAUGES[2] - 0.12 * 360) <= 1


def test_title_cards_recede_one_more_frame_per_part(deck, specs):
    cards = [i for i, s in enumerate(specs) if s["kind"] == "title_card"]
    for part, i in enumerate(cards, start=1):
        names = [m.name for m in _motifs(deck.slides[i])]
        frames = {x.split("-")[3] for x in names if x.startswith("motif-line-depth")}
        assert len(frames) == part, (i + 1, part)
        assert sum(x.startswith("motif-line-corridor") for x in names) >= 8
        assert "photo-slot" in [sh.name for sh in deck.slides[i].shapes]


def test_letterbox_slides_carry_the_scene_ruler(deck, specs):
    total = len(specs)
    for n, (s, sl) in enumerate(zip(specs, deck.slides), start=1):
        ticks = sorted((m for m in _motifs(sl, "motif-line-tick")), key=lambda m: m.begin_x)
        if s["kind"] != "letterbox":
            assert not ticks, n
            continue
        assert len(ticks) == total, n
        lit = [k for k, m in enumerate(ticks, start=1) if abs(m.end_y - m.begin_y) / PX > 20]
        assert lit == [n], (n, lit)
        for m in ticks:
            assert m.begin_y / PX >= 942, (n, "ruler must sit in the footer")
            assert m.begin_y / PX >= max(t[3] for t in _texts(sl) if t[1] < 940), (n, "ruler below all content text")


def test_open_frames_have_a_horizon_below_all_text(deck, specs):
    for n, (s, sl) in enumerate(zip(specs, deck.slides), start=1):
        if s["layout"] != "close":
            continue
        hz = [m for m in _motifs(sl, "motif-horizon")]
        assert len(hz) == 1, n
        assert hz[0].top / PX >= max(t[3] for t in _texts(sl)), n


def test_lists_hang_on_a_rail(deck, specs):
    for n, (s, sl) in enumerate(zip(specs, deck.slides), start=1):
        if s["layout"] != "list":
            continue
        assert len(_motifs(sl, "motif-line-rail")) == 1, n
        assert len(_motifs(sl, "motif-line-stop")) == len(s["body"]), n


LINE_CHARTS = ["pawc_curve", "funnel", "vis_rank", "intent_named", "mirror_sites", "top_pages", "concentration",
               "sov_gap", "dead_links", "schema"]


def test_quantities_are_drawn_as_lines_not_bars(tmp_path):
    charts.render_all(load(), tmp_path)
    bars = charts.LAST_DATA["bars"]
    for k in LINE_CHARTS:
        assert bars[k] == 0, k


def test_text_boxes_never_overlap(deck):
    from deck.measure import height
    for n, s in enumerate(deck.slides, start=1):
        boxes = []
        for sh in s.shapes:
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            x0, y0, x1, _ = _box(sh)
            need = sum(height("".join(r.text for r in p.runs), x1 - x0, p.runs[0].font.size.pt,
                              p.runs[0].font.name == "IBM Plex Mono", bool(p.runs[0].font.bold))
                       for p in sh.text_frame.paragraphs if p.runs)
            boxes.append((x0, y0, x1, y0 + need, sh.text_frame.text[:30]))
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                assert not (a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]), (n, a[4], b[4])


def test_list_items_fill_the_slide_at_a_readable_size(deck, specs):
    # Eye check 2026-10-11: lists at 18 pt left half of the slide empty.
    for n, (s, sl) in enumerate(zip(specs, deck.slides), start=1):
        if s["layout"] != "list":
            continue
        items = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text in s["body"]]
        assert min(sh.text_frame.paragraphs[0].runs[0].font.size.pt for sh in items) >= 20, n


def test_hero_gauges_carry_the_slide(deck):
    g = [sh for sh in deck.slides[2].shapes if sh.name.startswith("motif-gauge")]
    assert min(sh.width / PX for sh in g) >= 280
