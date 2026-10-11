import re

import pytest
from pptx import Presentation
from pptx.util import Emu

from deck import build, copy
from deck.facts import load

PX = 6350  # EMU per slide pixel (1920 px wide)


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    path = build.build(tmp_path_factory.mktemp("deck") / "favoured-banks.pptx")
    return Presentation(str(path))


def _texts(deck):
    for i, s in enumerate(deck.slides, start=1):
        for sh in s.shapes:
            if sh.has_text_frame:
                yield i, sh, sh.text_frame.text
        yield i, None, s.notes_slide.notes_text_frame.text


def test_slide_count_and_kinds(deck):
    assert len(deck.slides) == 33
    kinds = [s["kind"] for s in copy.slides(load())]
    assert [i + 1 for i, k in enumerate(kinds) if k == "title_card"] == [5, 10, 18, 27]


def test_no_dashes_or_hype(deck):
    banned = ["seamless", "revolutionary", "cutting edge", "next generation", "powerful", "leverage", "unlock",
              "game-changer", "elevate", "delve"]
    for i, _, t in _texts(deck):
        assert "—" not in t and "–" not in t, i
        low = t.lower()
        assert not [w for w in banned if w in low], (i, t)


def test_fonts_only_sans_and_mono(deck):
    for s in deck.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        assert r.font.name in ("Plus Jakarta Sans", "IBM Plex Mono"), r.text


def test_numbers_come_from_facts(deck):
    allowed = copy.allowed_numbers(load())
    for i, _, t in _texts(deck):
        t = t.replace("09.10.2026", "")
        t = re.sub(r"^\d+ / \d+$", "", t)
        for tok in re.findall(r"\d[\d,.]*\d|\d", t):
            assert tok.rstrip(".,") in allowed, (i, tok, t[:80])


def test_text_fits(deck):
    from deck.measure import height
    for i, sh, _ in _texts(deck):
        if sh is None:
            continue
        w_px, h_px = sh.width / PX, sh.height / PX
        need = 0.0
        for p in sh.text_frame.paragraphs:
            if not p.runs:
                continue
            r = p.runs[0]
            need += height("".join(x.text for x in p.runs), w_px, r.font.size.pt, r.font.name == "IBM Plex Mono", bool(r.font.bold))
        assert need <= h_px + 2, (i, sh.text_frame.text[:60], round(need), round(h_px))


def test_letterbox_slides_have_slate(deck):
    kinds = [s["kind"] for s in copy.slides(load())]
    for n, (k, s) in enumerate(zip(kinds, deck.slides), start=1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        has = any(t == f"{n} / {len(kinds)}" for t in texts)
        assert has == (k == "letterbox"), (n, k)


def test_content_stays_inside_the_safe_area(deck):
    for n, s in enumerate(deck.slides, start=1):
        for sh in s.shapes:
            if sh.name.startswith(("motif", "ground", "chrome", "photo-slot")):
                continue
            x0, y0 = sh.left / PX, sh.top / PX
            x1, y1 = x0 + sh.width / PX, y0 + sh.height / PX
            assert 100 <= x0 and x1 <= 1820 and 80 <= y0 and y1 <= 1000, (n, sh.name, [round(v) for v in (x0, y0, x1, y1)])


def test_footer_fits_the_slide(deck):
    for n, s in enumerate(deck.slides, start=1):
        for sh in s.shapes:
            if sh.name.startswith("chrome"):
                assert (sh.top + sh.height) / PX <= 1060, (n, sh.name)


def test_no_tracked_capitals_or_dot_chains(deck):
    """Audit 002: wide-tracked uppercase labels and middle-dot chains read as AI defaults."""
    for n, s in enumerate(deck.slides, start=1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    rpr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    assert rpr is None or int(rpr.get("spc", "0")) <= 0, (n, r.text)
                    letters = re.sub(r"[^A-Za-z]", "", r.text)
                    shouted = len(letters) >= 6 and letters.isupper() and " " in r.text.strip()
                    assert not shouted, (n, r.text)  # an all-caps label, not an acronym such as BTPN
                    assert "·" not in r.text, (n, r.text)


ROLES = ["For the SEO strategist", "For content and marketing", "For the web team"]


def test_recommendations_address_the_people_who_run_the_website():
    specs = copy.slides(load())
    titles = [s["title"] for s in specs]
    assert not [t for t in titles if t.startswith("For banks") or "GEO practitioners" in t]
    start = titles.index(ROLES[0])
    assert titles[start:start + 3] == ROLES
    for s in specs[start:start + 3]:
        for item in s["body"]:
            assert re.search(r"\d", item), (s["title"], item)  # every recommendation carries its evidence
