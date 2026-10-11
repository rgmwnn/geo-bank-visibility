"""Assemble the editable deck (deck/out/favoured-banks.pptx) from facts, charts and copy."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from deck import charts, copy, measure
from deck.facts import load
from deck.theme import C, MONO, SANS

ROOT = Path(__file__).resolve().parent
PX = 6350
W, H = 1920, 1080
LEFT = 120
OUT = ROOT / "out" / "favoured-banks.pptx"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def est_h(t: str, w: float, pt: float, mono: bool = False, bold: bool = False) -> float:
    return measure.height(t, w, pt, mono, bold)


def rect(slide, x, y, w, h, color, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def seg(slide, x0, y0, x1, y1, color, weight=1.0, name=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x0 * PX)), Emu(int(y0 * PX)),
                                    Emu(int(x1 * PX)), Emu(int(y1 * PX)))
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Pt(weight)
    if name:
        ln.name = name
    return ln


def hline(slide, x0, x1, y, color, weight=1.5, name=None):
    return seg(slide, x0, y, x1, y, color, weight, name)


def _png(path: Path, w_px: int, h_px: int, draw) -> Path:
    """Draw with matplotlib on a transparent canvas measured in slide pixels (y grows downwards)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    path.parent.mkdir(parents=True, exist_ok=True)
    fig = plt.figure(figsize=(w_px / 100, h_px / 100), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    draw(ax)
    ax.set_xlim(0, w_px)
    ax.set_ylim(h_px, 0)
    ax.axis("off")
    fig.savefig(path, transparent=True)
    plt.close(fig)
    return path


def _place(slide, path: Path, x, y, w, h, name):
    pic = slide.shapes.add_picture(str(path), Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    pic.name = name
    return pic


ART = ROOT / "out" / "charts"

# Line language (deck/design-philosophy.md). Every motif is named motif-*; none may cross text.
# Grounds (ground-*) sit under text by design: the photo fade and the chart panels.

FADE_W, FADE_RAMP = 1400, 360  # navy is fully opaque up to x = 1040, then fades into the photo


def fade_png(width: int = FADE_W, ramp: int = FADE_RAMP) -> Path:
    import numpy as np
    path = ART / f"ground-fade-{width}-{ramp}.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    alpha = np.clip((width - np.arange(width)) / ramp, 0, 1)
    img = np.zeros((H, width, 4), dtype=np.uint8)
    img[..., :3] = [int(C["bg"][i:i + 2], 16) for i in (1, 3, 5)]
    img[..., 3] = (np.tile(alpha, (H, 1)) * 255).astype(np.uint8)
    Image.fromarray(img, "RGBA").save(path)
    return path


def photo_slot(slide):
    """Full-bleed stand-in for the tower photo; Canva fills it after import (update_fill keeps it at the back)."""
    rect(slide, 0, 0, W, H, "#1D3A63", name="photo-slot")
    _place(slide, fade_png(), 0, 0, FADE_W, H, "ground-fade")


REGION = (1060, 0, W, H)  # the corridor stands over the photo, right of the text
WALL = (1440, 420, 1680, 640)  # its far wall
DEPTHS = [0.5, 0.75, 0.875, 0.9375]


def corridor(slide, frames: int):
    """A tesseract corridor over the photo: rays to the far wall, plus one receding frame per part of the deck."""
    rx0, ry0, rx1, ry1 = REGION
    x0, y0, x1, y1 = WALL

    def at(t):
        return rx0 + (x0 - rx0) * t, ry0 + (y0 - ry0) * t, rx1 + (x1 - rx1) * t, ry1 + (y1 - ry1) * t

    for k, (ax, ay, bx, by) in enumerate([(rx0, ry0, x0, y0), (rx1, ry0, x1, y0), (rx0, ry1, x0, y1), (rx1, ry1, x1, y1)]):
        seg(slide, ax, ay, bx, by, C["glow"], 1.0, f"motif-line-corridor-corner-{k}")
    for k in range(1, 5):
        u = k / 5
        seg(slide, rx0 + (rx1 - rx0) * u, ry0, x0 + (x1 - x0) * u, y0, C["glow"], 0.5, f"motif-line-corridor-top-{k}")
        seg(slide, rx0 + (rx1 - rx0) * u, ry1, x0 + (x1 - x0) * u, y1, C["glow"], 0.5, f"motif-line-corridor-floor-{k}")
    for k in range(1, 4):
        v = k / 4
        seg(slide, rx1, ry0 + (ry1 - ry0) * v, x1, y0 + (y1 - y0) * v, C["glow"], 0.5, f"motif-line-corridor-right-{k}")
    for i, t in enumerate(DEPTHS[:frames], start=1):
        a, b, c, d = at(t)
        lit = i == frames
        for side, (p, q, r_, s_) in zip("tblr", [(a, b, c, b), (a, d, c, d), (a, b, a, d), (c, b, c, d)]):
            seg(slide, p, q, r_, s_, C["text"] if lit else C["glow"], 1.5 if lit else 0.75, f"motif-line-depth-{i}-{side}")


BRACKET = 40


def panel(slide, box, pad: int = 32):
    """A framed evidence panel: a darker ground with HUD corner brackets."""
    x, y, w, h = box
    rect(slide, x, y, w, h, C["panel"], name="ground-panel")
    x0, y0, x1, y1 = x - 10, y - 10, x + w + 10, y + h + 10
    for corner, (cx, cy, dx, dy) in {"tl": (x0, y0, 1, 1), "tr": (x1, y0, -1, 1),
                                     "bl": (x0, y1, 1, -1), "br": (x1, y1, -1, -1)}.items():
        seg(slide, cx, cy, cx + dx * BRACKET, cy, C["glow"], 1.5, f"motif-bracket-{corner}-h")
        seg(slide, cx, cy, cx, cy + dy * BRACKET, C["glow"], 1.5, f"motif-bracket-{corner}-v")
    return x + pad, y + pad, w - 2 * pad, h - 2 * pad


def ruler(slide, n: int, total: int):
    """One tick per slide in the footer; the current slide is lit."""
    x0, step, y = 1200, 15, 994
    for k in range(1, total + 1):
        x = x0 + (k - 1) * step
        if k == n:
            seg(slide, x, y - 14, x, y + 14, C["text"], 2.0, f"motif-line-tick-{k:02d}")
        else:
            seg(slide, x, y - 6, x, y + 6, C["secondary"] if k < n else C["line"], 1.0, f"motif-line-tick-{k:02d}")


GAUGES: dict[int, float] = {}


def gauge(slide, i: int, value: float, x, y, d):
    """A ring gauge: the steel track is the whole, the white arc is the share (sweep = value x 360)."""
    import numpy as np
    sweep = value * 360
    GAUGES[i] = sweep

    def draw(ax):
        c, r = d / 2, d / 2 - 10
        t = np.linspace(0, 2 * np.pi, 400)
        ax.plot(c + r * np.cos(t), c + r * np.sin(t), color=C["line"], lw=6)
        a = np.radians(-90 + np.linspace(0, sweep, 200))
        ax.plot(c + r * np.cos(a), c + r * np.sin(a), color=C["text"], lw=14, solid_capstyle="butt")
        for k in range(12):  # hour ticks inside the track, like an instrument dial
            ang = np.radians(k * 30)
            ax.plot([c + (r - 22) * np.cos(ang), c + (r - 32) * np.cos(ang)],
                    [c + (r - 22) * np.sin(ang), c + (r - 32) * np.sin(ang)], color=C["line"], lw=1.4)

    path = _png(ART / f"motif-gauge-{i}.png", d, d, draw)
    return _place(slide, path, x, y, d, d, f"motif-gauge-{i}")


def horizon(slide, top: float = 900):
    """A planet's horizon along the bottom edge: the two engine threads bent into one arc."""
    import numpy as np
    h = H - top

    def draw(ax):
        R = 2600
        xs = np.linspace(0, W, 1200)
        for off, col in ((4, C["gpt"]), (18, C["gem"])):
            ax.plot(xs, off + R - np.sqrt(R ** 2 - (xs - W / 2) ** 2), color=col, lw=1.6, solid_capstyle="butt")

    return _place(slide, _png(ART / "motif-horizon.png", W, h, draw), 0, top, W, h, "motif-horizon")


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    """paras: list of (text, pt, color, font, bold, tracking) tuples, one paragraph each.
    The box grows to the measured height of its text, so nothing overflows."""
    h = max(h, sum(est_h(t, w, pt, font == MONO, bold) for t, pt, _, font, bold, _ in paras))
    tb = slide.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (t, pt, color, font, bold, track) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = t
        r.font.name = font
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.color.rgb = _rgb(color)
        if track:
            r._r.get_or_add_rPr().set("spc", str(track))
    return tb


def P(t, pt, color=C["text"], font=SANS, bold=False, track=0):
    return (t, pt, color, font, bold, track)


def picture(slide, path: Path, box, valign: str = "top"):
    x, y, w, h = box
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    top = y + (h - ph) / 2 if valign == "middle" else y
    pic = slide.shapes.add_picture(str(path), Emu(int((x + (w - pw) / 2) * PX)), Emu(int(top * PX)),
                                   Emu(int(pw * PX)), Emu(int(ph * PX)))
    pic.name = f"chart-{Path(path).stem}"
    return pic


def framed_chart(slide, path: Path, box, valign: str = "middle"):
    """Fit the chart, then draw the panel tight around it so the brackets frame the evidence, not empty space."""
    x, y, w, h = box
    pad = 14
    iw, ih = Image.open(path).size
    scale = min((w - 2 * pad) / iw, (h - 2 * pad) / ih)
    pw, ph = iw * scale + 2 * pad, ih * scale + 2 * pad
    py = y + (h - ph) / 2 if valign == "middle" else y
    px = x + (w - pw) / 2
    inner = panel(slide, (px, py, pw, ph), pad)
    return picture(slide, path, inner, valign="middle")


def frame(slide, kind: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(C["bg"])


FOOT_LINE, FOOT_Y, CONTENT_BOTTOM = 952, 970, 930


def chrome(slide, n: int, s: dict, total: int):
    hline(slide, LEFT, LEFT + 32, 86, C["glow"], 2.0, "motif-line-mark")
    text(slide, LEFT + 48, 72, 900, 28, [P(s.get("part", ""), 15, C["secondary"])], name="chrome-part")
    hline(slide, LEFT, RIGHT, FOOT_LINE, C["line"], 0.75, "motif-line-footer")
    text(slide, 1700, FOOT_Y + 10, 100, 26, [P(f"{n} / {total}", 10, C["secondary"], MONO)], align=PP_ALIGN.RIGHT,
         name="chrome-page")
    if s.get("foot"):
        text(slide, LEFT, FOOT_Y, 1040, 26, [P(s["foot"], 10, C["secondary"], MONO)], name="chrome-foot")


RIGHT = W - LEFT
TITLE_PT = 34


def _title(slide, s, x, y, w, pt=TITLE_PT):
    h = est_h(s["title"], w, pt, bold=True)
    text(slide, x, y, w, h, [P(s["title"], pt, bold=True)])
    return y + h


def _body(slide, paras, x, y, w, pt=15, gap=14, color=C["secondary"], bullets=False):
    for b in paras:
        t = ("•  " + b) if bullets else b
        h = est_h(t, w, pt)
        text(slide, x, y, w, h, [P(t, pt, color)])
        y += h + gap
    return y


TEXT_W = FADE_W - FADE_RAMP - LEFT - 40  # photo pages keep text on the opaque navy


def lay_cover(slide, s, n, art):
    photo_slot(slide)
    text(slide, LEFT, 150, TEXT_W, 30, [P(s["body"][1], 15, C["secondary"], MONO)])
    th = est_h(s["title"], TEXT_W, 80, bold=True)
    text(slide, LEFT, 196, TEXT_W, th, [P(s["title"], 80, bold=True)])
    y = 196 + th + 24
    rect(slide, LEFT + 4, y, 120, 6, C["gpt"], name="motif-rule-cover")
    y += 6 + 36
    sh = est_h(s["body"][0], TEXT_W, 24)
    text(slide, LEFT, y, TEXT_W, sh, [P(s["body"][0], 24, C["secondary"])])
    text(slide, LEFT, max(y + sh + 40, 930), TEXT_W, 30, [P(s["byline"], 15, C["text"], MONO)])


def lay_hero(slide, s, n, art):
    y = _title(slide, s, LEFT, 150, 1680)
    top = max(380, y + 60)
    d = 300
    for i, (big, cap, src) in enumerate(s["heroes"], start=1):
        x = LEFT + (i - 1) * 840
        gauge(slide, i, float(big.rstrip("%")) / 100, x, top, d)
        bh = est_h(big, 480, 100, bold=True)
        text(slide, x + d + 44, top + (d - bh) / 2, 480, bh, [P(big, 100, bold=True)])
        caph = est_h(cap, 720, 18)
        text(slide, x, top + d + 36, 720, caph, [P(cap, 18, C["secondary"])])
        text(slide, x, top + d + 36 + caph + 14, 720, 30, [P(src, 11, C["secondary"], MONO)])


CARD_PARTS = {"Part I": 1, "Part II": 2, "Part III": 3, "Part IV": 4}


def lay_card(slide, s, n, art):
    photo_slot(slide)
    corridor(slide, CARD_PARTS[s["part"]])
    ph = est_h(s["part"], TEXT_W, 18)
    th = est_h(s["title"], TEXT_W, 80, bold=True)
    bh = est_h(s["body"][0], TEXT_W, 24)
    y = (H - (ph + 40 + th + 30 + bh)) / 2
    text(slide, LEFT, y, TEXT_W, ph, [P(s["part"], 18, C["secondary"])])
    hline(slide, LEFT, LEFT + 48, y + ph + 20, C["glow"], 2.0, "motif-line-card-rule")
    y += ph + 40
    text(slide, LEFT, y, TEXT_W, th, [P(s["title"], 80, bold=True)])
    text(slide, LEFT, y + th + 30, TEXT_W, bh, [P(s["body"][0], 24, C["secondary"])])


SIDE_W = 580
SIDE_SIZES = [(34, 15), (30, 15), (30, 14), (28, 14), (28, 13)]


def _side_fit(s):
    """Largest title and body sizes whose text column ends above the footer."""
    gap = 10 if s.get("bullets") else 14
    for tpt, bpt in SIDE_SIZES:
        y = 150 + est_h(s["title"], SIDE_W, tpt, bold=True) + 28
        y += sum(est_h(("•  " + b) if s.get("bullets") else b, SIDE_W, bpt) + gap for b in s["body"])
        if y <= CONTENT_BOTTOM:
            break
    return tpt, bpt


def lay_side(slide, s, n, art):
    tpt, bpt = _side_fit(s)
    y = _title(slide, s, LEFT, 150, SIDE_W, pt=tpt)
    bullets = s.get("bullets", False)
    _body(slide, s["body"], LEFT, y + 28, SIDE_W, pt=bpt, gap=10 if bullets else 14, bullets=bullets)
    if s.get("chart"):
        framed_chart(slide, art[s["chart"]], (740, 140, RIGHT - 740, CONTENT_BOTTOM - 140))


def lay_stack(slide, s, n, art):
    y = _title(slide, s, LEFT, 140, 1680, pt=30)
    y = _body(slide, s["body"], LEFT, y + 12, 1400, gap=8)
    framed_chart(slide, art[s["chart"]], (LEFT, y + 10, RIGHT - LEFT, CONTENT_BOTTOM + 10 - (y + 10)))


RAIL_X = LEFT + 82
LIST_SIZES = [(24, 34), (23, 30), (22, 26), (21, 22), (20, 18), (20, 12)]


def lay_list(slide, s, n, art):
    y0 = _title(slide, s, LEFT, 150, 1680) + 32
    numbered = s.get("numbered", False)
    x = LEFT + 120
    w = RIGHT - x
    for pt, gap in LIST_SIZES:
        if y0 + sum(est_h(i, w, pt) + gap for i in s["body"]) <= CONTENT_BOTTOM:
            break
    y = y0
    stops = []
    for i, item in enumerate(s["body"], start=1):
        h = est_h(item, w, pt)
        if numbered:
            text(slide, LEFT, y + pt * 0.35, 60, 34, [P(f"{i:02d}", 15, C["glow"], MONO)])
        text(slide, x, y, w, h, [P(item, pt)])
        stops.append(y + pt * measure.LINE)
        y += h + gap
    seg(slide, RAIL_X, stops[0] - 18, RAIL_X, stops[-1] + 18, C["line"], 1.0, "motif-line-rail")
    for k, sy in enumerate(stops, start=1):
        seg(slide, RAIL_X - 7, sy, RAIL_X + 7, sy, C["glow"], 1.5, f"motif-line-stop-{k}")


def lay_metrics(slide, s, n, art):
    y = _title(slide, s, LEFT, 150, 1680) + 24
    for term, definition in s["body"]:
        th = est_h(term, 800, 15, bold=True)
        text(slide, LEFT, y, 800, th, [P(term, 15, bold=True)])
        h = est_h(definition, 800, 13)
        text(slide, LEFT, y + th + 2, 800, h, [P(definition, 13, C["secondary"])])
        y += th + 2 + h + 12
    framed_chart(slide, art[s["chart"]], (960, 240, RIGHT - 960, 620))


def lay_cold_open(slide, s, n, art):
    y = _title(slide, s, LEFT, 150, 1680) + 34
    prompt = f"“{s['prompt']}”"
    h = est_h(prompt, 1500, 24)
    text(slide, LEFT, y, 1500, h, [P(prompt, 24)])
    y += h + 24
    h = est_h(s["english"], 1500, 15)
    text(slide, LEFT, y, 1500, h, [P(s["english"], 15, C["secondary"])])
    y += h + 60
    for i, (head, banks, key) in enumerate(s["cols"]):
        x = LEFT + i * 760
        rect(slide, x, y + 8, 14, 14, C[key])
        text(slide, x + 30, y - 2, 600, 30, [P(head, 15, C["secondary"])])
        names = banks.split("\n")
        text(slide, x, y + 50, 680, len(names) * 62, [P(b, 26) for b in names])


def lay_close(slide, s, n, art):
    th = est_h(s["title"], 1680, 40, bold=True)
    text(slide, LEFT, 150, 1680, th, [P(s["title"], 40, bold=True)])
    text(slide, LEFT, 150 + th + 14, 1600, 40, [P(s["body"][0], 17, C["text"], MONO)])
    y = max(390, 150 + th + 14 + est_h(s["body"][0], 1600, 17, mono=True) + 50)
    text(slide, LEFT, y, 400, 28, [P("Sources", 15, C["secondary"])])
    y += 46
    for src in s["sources"]:
        text(slide, LEFT, y, 1600, 28, [P(src, 12, C["secondary"], MONO)])
        y += 38
    text(slide, LEFT, max(800, y + 24), 1300, 34, [P(s["byline"], 14, C["text"], MONO)])
    horizon(slide)


LAYOUTS = {"cover": lay_cover, "hero": lay_hero, "card": lay_card, "side": lay_side, "stack": lay_stack,
           "list": lay_list, "metrics": lay_metrics, "cold_open": lay_cold_open, "close": lay_close}


def build(out: Path = OUT) -> Path:
    f = load()
    art = charts.render_all(f, ART)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    blank = prs.slide_layouts[6]
    specs = copy.slides(f)
    for n, s in enumerate(specs, start=1):
        slide = prs.slides.add_slide(blank)
        frame(slide, s["kind"])
        LAYOUTS[s["layout"]](slide, s, n, art)
        if s["kind"] == "letterbox":
            chrome(slide, n, s, len(specs))
            ruler(slide, n, len(specs))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    print(build())
