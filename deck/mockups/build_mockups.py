"""Spike: three dark-blue design directions, two slides each (cover + one finding). Throwaway."""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

from deck.facts import load  # noqa: E402

OUT = Path(__file__).parent / "out"
OUT.mkdir(exist_ok=True)
PX, W, H = 6350, 1920, 1080

DIRS = {
    "A": dict(name="Annual report", bg="#F6F3EC", ink="#13294B", sub="#4A5568", navy="#13294B", on_navy="#F6F3EC",
              on_navy_sub="#B9C6DA", accent="#8F6212", gpt="#B4651C", gem="#2E6CB5",
              head="Instrument Serif", body="Instrument Sans", head_bold=False),
    "B": dict(name="Consulting", bg="#FFFFFF", ink="#0E2A52", sub="#4B5563", navy="#0E2A52", on_navy="#FFFFFF",
              on_navy_sub="#C9D4E5", accent="#D2622F", gpt="#D2622F", gem="#2A64B4",
              head="Bricolage Grotesque", body="Instrument Sans", head_bold=True),
    "C": dict(name="Navy night", bg="#0F2747", ink="#F4F6FA", sub="#AFC0D8", navy="#0F2747", on_navy="#F4F6FA",
              on_navy_sub="#AFC0D8", accent="#F2A541", gpt="#F2A541", gem="#7FC0F5",
              head="Outfit", body="Outfit", head_bold=True),
}


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def rect(s, x, y, w, h, color, name=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(color)
    r.line.fill.background()
    r.shadow.inherit = False
    if name:
        r.name = name
    return r


FILES = {("Instrument Serif", False, False): "InstrumentSerif-Regular.ttf",
         ("Instrument Serif", False, True): "InstrumentSerif-Italic.ttf",
         ("Instrument Sans", False, False): "InstrumentSans-Regular.ttf",
         ("Instrument Sans", True, False): "InstrumentSans-Bold.ttf",
         ("Bricolage Grotesque", True, False): "BricolageGrotesque-Bold.ttf",
         ("Bricolage Grotesque", False, False): "BricolageGrotesque-Regular.ttf",
         ("Outfit", True, False): "Outfit-Bold.ttf", ("Outfit", False, False): "Outfit-Regular.ttf"}


def measure(t, w, size, font, bold=False, italic=False):
    from PIL import ImageFont
    f = ImageFont.truetype(str(Path("~/.fonts").expanduser() / FILES[(font, bold, italic)]), int(size * 2))
    lines, cur = 1, ""
    for word in t.split(" "):
        trial = f"{cur} {word}".strip()
        if cur and f.getlength(trial) > w * 0.94:
            lines, cur = lines + 1, word
        else:
            cur = trial
    return lines * size * 2 * 1.42


def text(s, x, y, w, h, t, size, color, font, bold=False, italic=False, align=PP_ALIGN.LEFT, spacing=1.15):
    h = max(h, measure(t, w, size, font, bold, italic))
    tb = s.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int(h * PX)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = t
    r.font.size, r.font.name, r.font.bold, r.font.italic = Pt(size), font, bold, italic
    r.font.color.rgb = rgb(color)
    tb.bottom = y + h
    return tb


def chart(d, f, path):
    from matplotlib import font_manager
    for ttf in Path("~/.fonts").expanduser().glob("*.ttf"):
        font_manager.fontManager.addfont(str(ttf))
    banks = list(f["vis"].sort_values("n_answers", ascending=False).brand.head(10))
    banks = sorted(banks, key=lambda b: -f["mention_rate"][b].get("All", 0))[:10]
    mr = f["mention_rate"]
    plt.rcParams.update({"font.family": d["body"], "font.size": 14})
    fig, ax = plt.subplots(figsize=(7.4, 5.2), dpi=260)
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")
    fig.subplots_adjust(left=0.31, right=0.97, top=0.86, bottom=0.03)
    n = len(banks)
    for t in (0, 0.2, 0.4, 0.6):
        ax.plot([t, t], [-0.6, n - 0.4], color=d["sub"], lw=0.6, alpha=0.35, zorder=0)
        ax.text(t, n - 0.15, f"{int(t * 100)}%", ha="center", fontsize=14, color=d["sub"])
    for i, b in enumerate(banks):
        y = n - 1 - i
        g, m = mr[b].get("ChatGPT", 0), mr[b].get("Gemini", 0)
        lit = i < 3
        ax.text(-0.02, y, b, ha="right", va="center", fontsize=15, color=d["ink"] if lit else d["sub"],
                fontweight="bold" if lit else "normal")
        ax.plot([g, m], [y, y], color=d["sub"], lw=2.2, alpha=0.5, zorder=1)
        ax.scatter([g], [y], s=170, color=d["gpt"], zorder=3, edgecolor=d["bg"], linewidth=1.5)
        ax.scatter([m], [y], s=170, color=d["gem"], zorder=3, edgecolor=d["bg"], linewidth=1.5)
    for x, e, c in ((0.0, "ChatGPT", d["gpt"]), (0.2, "Gemini", d["gem"])):
        ax.scatter([x], [n + 0.85], s=150, color=c, clip_on=False)
        ax.text(x + 0.02, n + 0.85, e, va="center", fontsize=16, color=d["ink"])
    ax.set_xlim(-0.01, 0.7)
    ax.set_ylim(-0.6, n + 0.3)
    ax.axis("off")
    fig.savefig(path, transparent=True)
    plt.close(fig)
    return path


def fade(d, path):
    """Navy panel that fades into the photo on its right edge (direction C)."""
    w, h = 1200, 1080
    a = np.ones((h, w))
    ramp = np.clip((w - np.arange(w)) / 360, 0, 1)
    a = a * np.minimum(1, ramp)[None, :]
    col = np.array([int(d["navy"][i:i + 2], 16) for i in (1, 3, 5)], dtype=np.uint8)
    img = np.zeros((h, w, 4), dtype=np.uint8)
    img[..., :3] = col
    img[..., 3] = (a * 255).astype(np.uint8)
    Image.fromarray(img, "RGBA").save(path)
    return path


def cover(prs, key, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    title, sub = "Favoured Banks", "Which Indonesian banks ChatGPT and Gemini recommend, and which web pages they cite"
    meta, by = "40 answers, collected on 9 October 2026", "Rahmananda Ridho Gamawan, data analyst"
    if key == "A":
        s.background.fill.fore_color.rgb = rgb(d["navy"])
        rect(s, 1110, 0, 810, H, "#1D3A63", name="photo-slot")  # the tower photo lands here in Canva
        text(s, 120, 250, 900, 40, meta, 16, d["on_navy_sub"], d["body"])
        y = text(s, 120, 300, 1010, 10, title, 84, d["on_navy"], d["head"]).bottom
        rect(s, 124, y + 10, 140, 4, "#C9A65A")
        y = text(s, 120, y + 50, 900, 10, sub, 28, d["on_navy_sub"], d["head"], italic=True).bottom
        text(s, 120, max(y + 60, 930), 900, 40, by, 15, d["on_navy"], d["body"])
    elif key == "B":
        s.background.fill.fore_color.rgb = rgb(d["bg"])
        rect(s, 0, 0, W, 560, "#D9E1EC", name="photo-slot")
        rect(s, 0, 560, W, 14, d["navy"])
        y = text(s, 120, 610, 1680, 10, title, 76, d["ink"], d["head"], bold=True).bottom
        y = text(s, 120, y + 6, 1400, 10, sub, 26, d["sub"], d["body"]).bottom
        text(s, 120, max(y + 30, 960), 1680, 40, f"{by}. {meta}", 15, d["ink"], d["body"])
    else:
        s.background.fill.fore_color.rgb = rgb("#1D3A63")
        rect(s, 0, 0, W, H, "#1D3A63", name="photo-slot")
        pic = s.shapes.add_picture(str(fade(d, OUT / "fade_C.png")), 0, 0, Emu(1200 * PX), Emu(H * PX))
        pic.name = "fade"
        text(s, 120, 180, 900, 40, meta, 16, d["on_navy_sub"], d["body"])
        y = text(s, 120, 230, 1000, 10, title, 72, d["on_navy"], d["head"], bold=True).bottom
        rect(s, 124, y + 14, 120, 6, d["accent"])
        y = text(s, 120, y + 56, 820, 10, sub, 24, d["on_navy_sub"], d["body"]).bottom
        text(s, 120, max(y + 60, 930), 900, 40, by, 15, d["on_navy"], d["body"])
    return s


def finding(prs, key, d, f):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(d["bg"])
    png = chart(d, f, OUT / f"chart_{key}.png")
    title = "BCA, Bank Jago and SeaBank are each named in more than half of all answers"
    body = "BCA and Bank Jago appear in 22 of 40 answers, SeaBank in 21 and BRI in 12."
    part, foot = "Part II: The answers", "Source: data/processed/brand_engine.csv. Top 10 banks, n = 20 answers per engine"
    top = 150
    if key == "B":
        title = "BCA, Bank Jago and SeaBank are each named in over half of all answers"
        band = measure(title, 1680, 34, d["head"], True) + 130
        rect(s, 0, 0, W, band, d["navy"])
        text(s, 120, 40, 1680, 40, part, 15, d["on_navy_sub"], d["body"])
        text(s, 120, 82, 1680, 10, title, 34, d["on_navy"], d["head"], bold=True)
        top = band + 40
        text(s, 120, top + 40, 560, 10, body, 22, d["sub"], d["body"])
    else:
        if key == "A":
            rect(s, 120, 96, 60, 4, d["accent"])
        text(s, 120, 120, 640, 40, part, 15, d["sub"], d["body"])
        y = text(s, 120, 170, 640, 10, title, 42 if key == "A" else 36, d["ink"], d["head"], bold=d["head_bold"]).bottom
        text(s, 120, y + 30, 600, 10, body, 22, d["sub"], d["body"])
    iw, ih = Image.open(png).size
    bw, bh = 1060, min(760, 975 - (top + 10))
    sc = min(bw / iw, bh / ih)
    s.shapes.add_picture(str(png), Emu(int(780 * PX)), Emu(int((top + 10) * PX)), Emu(int(iw * sc * PX)),
                         Emu(int(ih * sc * PX)))
    text(s, 120, 1000, 1400, 30, foot, 11, d["sub"], d["body"])
    text(s, 1700, 1000, 100, 30, "11 / 33", 11, d["sub"], d["body"], align=PP_ALIGN.RIGHT)
    return s


def main():
    f = load()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    for key, d in DIRS.items():
        cover(prs, key, d)
        finding(prs, key, d, f)
    out = OUT / "mockups.pptx"
    prs.save(str(out))
    print(out)


if __name__ == "__main__":
    main()
